"""Build the editable two-slide EgoVerse interim PowerPoint.

Run from the repository root with:
source emimic/bin/activate
PYTHONPATH=.local_calibration/python_deps /usr/bin/python3 \
    docs/build_egoverse_reproduction_pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("egoverse_reproduction_update.pptx")

INK = "132238"
MUTED = "607084"
PAPER = "F6F3ED"
PANEL = "FFFDFA"
LINE = "D9D6CF"
ROBOT = "008B92"
ROBOT_SOFT = "DEF0F1"
HUMAN = "E47C24"
HUMAN_SOFT = "FAE8D8"
SHARED = "6555B5"
SHARED_SOFT = "E9E5F7"
GREEN = "27815E"
GREEN_SOFT = "E2F0E9"
FM_A = "506982"
FM_B = "93A6B7"
MLP_A = "DC741E"
MLP_B = "F3B27C"
FONT = "Microsoft YaHei"


def rgb(value):
    return RGBColor.from_string(value)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_line(shape, color=LINE, width=1):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=12,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rich_text(slide, x, y, w, h, lines, margin=0.06, valign=MSO_ANCHOR.TOP):
    """Add editable paragraphs; line = (text, size, color, bold, spacing_after)."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, (text, size, color, bold, after) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(after)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def add_box(slide, x, y, w, h, accent, title, body_lines, fill=PANEL):
    body = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(body, fill)
    set_line(body)
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h)
    )
    set_fill(accent_bar, accent)
    accent_bar.line.fill.background()
    add_text(slide, x + 0.15, y + 0.1, w - 0.25, 0.25, title, 12, accent, True)
    add_rich_text(slide, x + 0.13, y + 0.39, w - 0.23, h - 0.45, body_lines, 0.02)
    return body


def add_arrow(slide, x1, y1, x2, y2):
    arrow = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    arrow.line.color.rgb = rgb(SHARED)
    arrow.line.width = Pt(2)
    line = arrow.line._get_or_add_ln()
    arrowhead = OxmlElement("a:headEnd")
    arrowhead.set("type", "triangle")
    line.append(arrowhead)
    return arrow


def add_header(slide, eyebrow, title, subtitle, number):
    set_fill(slide.background, PAPER)
    for x, w, color in [(0, 4.45, ROBOT), (4.45, 4.45, SHARED), (8.9, 4.433, HUMAN)]:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), 0, Inches(w), Inches(0.07)
        )
        set_fill(bar, color)
        bar.line.fill.background()
    add_text(slide, 0.5, 0.2, 8.0, 0.2, eyebrow.upper(), 9, SHARED, True)
    add_text(slide, 0.5, 0.45, 12.2, 0.48, title, 26, INK, True)
    add_text(slide, 0.5, 0.93, 11.8, 0.28, subtitle, 11, MUTED)
    add_text(
        slide, 12.4, 7.12, 0.45, 0.18, f"{number} / 2", 8, MUTED, False, PP_ALIGN.RIGHT
    )


def add_loss_card(slide, x, w, title, description, equation):
    y, h = 5.55, 1.28
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(card, PANEL)
    set_line(card)
    add_text(slide, x + 0.12, y + 0.09, w - 0.24, 0.22, title, 10, INK, True)
    add_text(slide, x + 0.12, y + 0.34, w - 0.24, 0.38, description, 7.5, MUTED)
    add_text(slide, x + 0.12, y + 0.82, w - 0.24, 0.26, equation, 10, INK, False)


def architecture_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "EgoVerse reproduction · current policy",
        "Model Architecture｜双域、共享 HPT、双动作头",
        "A/B × MLP / Flow Matching · action horizon = 30",
        1,
    )

    columns = [
        (0.5, 2.15, "1 · INPUT"),
        (2.92, 1.62, "2 · ENCODER"),
        (4.82, 1.8, "3 · HPT"),
        (6.9, 2.15, "4 · ACTION HEADS"),
        (9.33, 2.15, "5 · OUTPUT"),
    ]
    for x, w, title in columns:
        add_text(slide, x, 1.28, w, 0.2, title, 9, MUTED, True, PP_ALIGN.CENTER)

    add_box(
        slide,
        0.5,
        1.55,
        2.15,
        1.7,
        ROBOT,
        "Piper domain",
        [
            ("Front RGB + Wrist RGB", 8.5, INK, True, 1),
            ("2 × [B, 3, 224, 224]", 8.5, MUTED, False, 4),
            ("Hardware state", 8.5, INK, True, 1),
            ("[q₁…q₆, gripper] = [B, 7]", 8.2, MUTED, False, 0),
        ],
    )
    add_box(
        slide,
        0.5,
        3.36,
        2.15,
        1.7,
        HUMAN,
        "Aria domain · B only",
        [
            ("Egocentric RGB", 8.5, INK, True, 1),
            ("[B, 3, 224, 224]", 8.5, MUTED, False, 4),
            ("Left-hand EE pose", 8.5, INK, True, 1),
            ("[xyz, yaw/pitch/roll] = [B, 6]", 8.0, MUTED, False, 0),
        ],
    )
    add_box(
        slide,
        2.92,
        1.55,
        1.62,
        3.51,
        SHARED,
        "Domain-aware stems",
        [
            ("ResNet18 / image", 8.5, INK, True, 1),
            ("[B, 49, 256]", 8.5, MUTED, False, 5),
            ("State MLP", 8.5, INK, True, 1),
            ("[B, 1, 256]", 8.5, MUTED, False, 5),
            ("Cross Attention", 8.5, INK, True, 1),
            ("each modality → [B, 16, 256]", 7.7, MUTED, False, 7),
            (
                "主视角 encoder 共享；腕部与 proprio stem 为 Piper 专用。",
                7.4,
                MUTED,
                False,
                0,
            ),
        ],
    )
    add_box(
        slide,
        4.82,
        1.55,
        1.8,
        3.51,
        SHARED,
        "Shared HPT trunk",
        [
            ("16 Transformer blocks", 8.5, INK, True, 1),
            ("width 256 · 8 heads", 8.2, MUTED, False, 5),
            ("Piper tokens", 8.5, INK, True, 1),
            ("30 + 3×16 = [B, 78, 256]", 7.8, MUTED, False, 5),
            ("Human tokens", 8.5, INK, True, 1),
            ("30 + 2×16 = [B, 62, 256]", 7.8, MUTED, False, 5),
            ("Action features", 8.5, INK, True, 1),
            ("h = [B, 30, 256]", 8.2, MUTED, False, 5),
            ("B: shared trunk + domain embedding", 7.5, MUTED, False, 0),
        ],
    )
    add_box(
        slide,
        6.9,
        1.55,
        2.15,
        1.7,
        SHARED,
        "Cartesian pose head",
        [
            ("Piper D=7 · Human D=6", 8.2, MUTED, False, 5),
            ("MLP: deterministic chunk", 8.2, MLP_A, True, 2),
            ("FM: 6-block CrossTransformer", 8.0, FM_A, True, 1),
            ("128D · 4 heads", 7.8, MUTED, False, 0),
        ],
    )
    add_box(
        slide,
        6.9,
        3.36,
        2.15,
        1.7,
        ROBOT,
        "Piper joint head",
        [
            ("Piper only · D=7", 8.2, MUTED, False, 5),
            ("MLP: hidden 256", 8.2, MLP_A, True, 2),
            ("FM: 50 inference steps", 8.0, FM_A, True, 2),
            ("同一 HPT feature，独立监督", 7.8, MUTED, False, 0),
        ],
    )
    add_box(
        slide,
        9.33,
        1.55,
        2.15,
        1.7,
        GREEN,
        "Geometry supervision",
        [
            ("Robot actions_cartesian", 8.3, INK, True, 1),
            ("[B, 30, 7]", 8.2, MUTED, False, 4),
            ("Human actions_cartesian", 8.3, INK, True, 1),
            ("[B, 30, 6]", 8.2, MUTED, False, 4),
            ("xyz + yaw/pitch/roll (+ gripper)", 7.5, MUTED, False, 0),
        ],
    )
    add_box(
        slide,
        9.33,
        3.36,
        2.15,
        1.7,
        GREEN,
        "Robot deployment",
        [
            ("actions_joints [B, 30, 7]", 8.2, INK, True, 4),
            ("6 absolute joint targets", 8.0, MUTED, False, 1),
            ("+ continuous gripper opening", 8.0, MUTED, False, 5),
            ("Direct joint control · no IK", 8.3, GREEN, True, 0),
        ],
    )
    for x1, x2 in [(2.68, 2.88), (4.57, 4.78), (6.65, 6.86), (9.08, 9.29)]:
        add_arrow(slide, x1, 3.28, x2, 3.28)

    add_loss_card(
        slide,
        0.5,
        2.65,
        "MLP head loss",
        "Normalized chunk MSE；无额外 rotation geodesic 项。",
        "L_MLP = MSE(a_hat, a)",
    )
    add_loss_card(
        slide,
        3.27,
        3.35,
        "Flow Matching head loss",
        "t~Beta(1.5,1); x_t=t eps+(1−t)a; target u_t=eps−a",
        "L_FM = E ||v_theta(x_t,t,h) − u_t||²",
    )
    add_loss_card(
        slide,
        6.74,
        2.1,
        "A · robot only",
        "Pose 与 joint 两个 head 等权相加。",
        "L_A = L_r,pose + L_r,joint",
    )
    add_loss_card(
        slide,
        8.96,
        3.87,
        "B · robot + human",
        "每 step 各取一个 domain batch；最终对两域求均值。",
        "L_B = ½(L_r,pose + L_r,joint) + ½L_h,pose",
    )
    return slide


def config_card(slide, x, w, accent, title, lines):
    y, h = 1.38, 1.03
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(card, PANEL)
    card.line.fill.background()
    top = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.05)
    )
    set_fill(top, accent)
    top.line.fill.background()
    add_text(slide, x + 0.12, y + 0.12, w - 0.24, 0.2, title, 10, INK, True)
    add_text(slide, x + 0.12, y + 0.36, w - 0.24, 0.56, lines, 7.7, MUTED)


def results_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Training & offline evaluation",
        "Training Configuration｜训练设置与离线结果",
        "固定 robot validation split：7 episodes · 2,381 frames",
        2,
    )
    config_card(
        slide,
        0.5,
        3.0,
        ROBOT,
        "Dataset",
        "A: 63 Piper episodes · 22,406 frames\nB: A + 45 Aria episodes · 7,546 frames\nSplit seed = 42",
    )
    config_card(
        slide,
        3.62,
        2.65,
        HUMAN,
        "Optimization",
        "10,000 updates · batch 8/domain\nAdamW · lr 1e−4 · wd 1e−4\ncosine → 1e−5 · BF16",
    )
    config_card(
        slide,
        6.39,
        2.65,
        SHARED,
        "Sequence",
        "30 Hz · observation horizon 1\naction horizon 30 = 1.0 s\nRGB 224² + ColorJitter",
    )
    config_card(
        slide,
        9.16,
        3.67,
        SHARED,
        "Compared variants",
        "A+MLP · B+MLP · A+FM · B+FM\nsame HPT trunk and robot split\nFM uses fixed inference seeds",
    )

    add_text(
        slide,
        0.5,
        2.55,
        7.0,
        0.25,
        "Offline Cartesian benchmark（越低越好）",
        13,
        INK,
        True,
    )
    chart_data = ChartData()
    chart_data.categories = [
        "First-step TCP\nposition",
        "30-step TCP\nposition",
        "Adjacent action\ndisplacement",
    ]
    chart_data.add_series("FM A", (104.73, 97.12, 66.19))
    chart_data.add_series("FM B", (107.62, 101.00, 65.44))
    chart_data.add_series("MLP A", (30.93, 49.63, 13.05))
    chart_data.add_series("MLP B", (36.58, 54.17, 12.93))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(2.85),
        Inches(9.55),
        Inches(3.63),
        chart_data,
    ).chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 120
    chart.value_axis.major_unit = 20
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.size = Pt(8)
    for series, color in zip(chart.series, (FM_A, FM_B, MLP_A, MLP_B)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color)
        series.format.line.fill.background()
        series.has_data_labels = True
        labels = series.data_labels
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
        labels.font.name = FONT
        labels.font.size = Pt(7)
        labels.font.color.rgb = rgb(INK)
        labels.number_format = "0.0"
    chart.legend.font.name = FONT
    chart.legend.font.size = Pt(8)

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.25),
        Inches(2.55),
        Inches(2.58),
        Inches(3.93),
    )
    set_fill(panel, PANEL)
    set_line(panel)
    add_text(slide, 10.43, 2.73, 2.2, 0.25, "Additional metrics", 12, INK, True)
    rows = [
        ("Model", "Rot. °", "Grip. mm"),
        ("FM A", "18.44", "15.35"),
        ("FM B", "20.57", "14.32"),
        ("MLP A", "11.06", "5.24"),
        ("MLP B", "12.61", "4.95"),
    ]
    table_shape = slide.shapes.add_table(
        len(rows), 3, Inches(10.42), Inches(3.08), Inches(2.22), Inches(1.62)
    )
    table = table_shape.table
    table.columns[0].width = Inches(0.86)
    table.columns[1].width = Inches(0.62)
    table.columns[2].width = Inches(0.74)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.margin_left = cell.margin_right = Inches(0.03)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(SHARED_SOFT if r_idx == 0 else PANEL)
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.RIGHT
            run = paragraph.runs[0]
            run.font.name = FONT
            run.font.size = Pt(7.5)
            run.font.bold = r_idx == 0
            run.font.color.rgb = rgb(MUTED if r_idx == 0 else INK)
    add_text(
        slide,
        10.43,
        4.8,
        2.18,
        0.35,
        "Rot. = 30-step orientation error\nGrip. = 30-step gripper MAE",
        7,
        MUTED,
    )
    note = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.42),
        Inches(5.28),
        Inches(2.22),
        Inches(0.93),
    )
    set_fill(note, ROBOT_SOFT)
    note.line.fill.background()
    add_rich_text(
        slide,
        10.52,
        5.37,
        2.02,
        0.72,
        [
            ("Observed trend", 8.5, ROBOT, True, 2),
            (
                "MLP显著降低位置误差与 chunk 内跳动；B 在当前 Piper 原场景验证集上未优于 A。",
                7.2,
                MUTED,
                False,
                0,
            ),
        ],
        0,
    )
    add_text(
        slide,
        0.5,
        6.62,
        12.1,
        0.32,
        "Cartesian pose-head完整离线评估；FM为原监控best（A 9.1k / B 9.8k），MLP为10k final。用于阶段性比较，不视为严格单变量因果实验；离线误差不能替代真机成功率。",
        7.2,
        MUTED,
    )
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    architecture_slide(prs)
    results_slide(prs)
    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
